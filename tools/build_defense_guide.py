from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

OUT = r"D:\Oranges\docs\OrangeOS答辩演示与问答手册.docx"
BLUE = "2E74B5"
DARK = "1F4D78"
LIGHT = "E8EEF5"
PALE = "F4F6F9"
GRAY = "5B6573"
RED = "9B1C1C"


def font(run, size=11, bold=False, color="000000", name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margin(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margin(row.cells[i])
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    tblW.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tblW.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    set_repeat_table_header(t.rows[0])
    for i, h in enumerate(headers):
        shade(t.rows[0].cells[i], LIGHT)
        p = t.rows[0].cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(h), 10, True, DARK)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            font(p.add_run(value), 9.5)
    set_table_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(text), 10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    font(p.add_run(text), 10.5)
    return p


def add_callout(doc, label, text, color=DARK):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    set_table_widths(t, [6.5])
    shade(t.cell(0, 0), PALE)
    p = t.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(label + "  "), 10.5, True, color)
    font(p.add_run(text), 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_q(doc, q, a, points=None):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run("问：" + q), 11, True, DARK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run("答：" + a), 10.5)
    if points:
        for x in points:
            add_bullet(doc, x, 1)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for name, size, before, after, color in [("Heading 1",16,18,10,BLUE),("Heading 2",13,14,7,BLUE),("Heading 3",12,10,5,DARK)]:
    s = styles[name]; s.font.name = "Microsoft YaHei"; s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after); s.paragraph_format.keep_with_next = True
for name in ["List Bullet", "List Bullet 2", "List Number"]:
    s=styles[name]; s.font.name="Microsoft YaHei"; s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    s.font.size=Pt(10.5); s.paragraph_format.space_after=Pt(4); s.paragraph_format.line_spacing=1.25

# Header/footer
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(hp.add_run("OrangeOS v1.0｜答辩准备"), 8.5, False, GRAY)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(fp.add_run("OrangeOS 答辩演示与问答手册"), 8.5, False, GRAY)

# Cover
for _ in range(4): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("ORANGEOS v1.0"), 14, True, BLUE)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(10)
font(p.add_run("答辩演示与问答手册"), 28, True, DARK)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("从 BIOS 启动到 Ring 3、OrangeFS 与交互式桌面"), 14, False, GRAY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(36)
font(p.add_run("用途：现场演示路线｜讲解话术｜高频问题｜故障兜底"), 11, True, BLUE)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(80)
font(p.add_run("项目目录：D:\\Oranges\n版本依据：README、DEMO、RELEASE_CHECKLIST 与 Makefile"), 10, False, GRAY)
doc.add_page_break()

add_heading(doc, "一、答辩的核心叙事", 1)
add_callout(doc, "一句话定位", "OrangeOS 是一个从 Boot Sector 和 Loader 开始自主实现、运行于 x86 32 位保护模式的教学型操作系统；它把启动、中断、调度、分页、用户态、系统调用、文件系统与交互界面串成了可运行、可验证的完整闭环。")
add_heading(doc, "建议始终围绕三个关键词", 2)
add_bullet(doc, "完整链路：BIOS → Boot Sector → Loader → 保护模式 → Kernel → Shell/GUI。")
add_bullet(doc, "机制可验证：不仅显示界面，还能观察调度、页分配、Ring 3、系统调用、异常隔离与资源回收。")
add_bullet(doc, "工程可回归：Makefile 中有 16 项 QEMU 自动测试，发布门禁为 make test-release。")
add_heading(doc, "开场话术（约 40 秒）", 2)
add_callout(doc, "可直接说", "各位老师好，我的项目是 OrangeOS。它是一个面向 x86 架构、从引导扇区开始自主实现的小型操作系统。项目重点不是堆叠界面，而是打通操作系统的关键机制：进入 32 位保护模式，建立 GDT、IDT 和分页，通过 PIT 中断进行抢占式调度，提供 Ring 3 用户态与 int 0x80 系统调用，并实现带冗余目录和数据校验的 OrangeFS。接下来我会按‘启动—内核机制—用户态隔离—文件系统—交互界面’这条主线演示。")

add_heading(doc, "二、现场准备与发布门禁", 1)
add_table(doc, ["阶段", "操作", "通过标准"], [
    ("答辩前一天", "在 WSL 项目目录执行 make test-release", "最终显示 PASSED (16/16)"),
    ("答辩前", "备份已验证的 orange.img 和对应提交", "现场故障可快速回退"),
    ("开场前", "提前启动 QEMU，确认键盘焦点与分辨率", "出现 OrangeOS>，硬件光标正常"),
    ("持久化演示前", "不要 make clean，不要使用 -snapshot", "重启后文件内容仍可读取"),
    ("结束操作", "fault、shutdown 只放到最后", "避免中途停机或关闭 QEMU"),
], [1.1, 2.7, 2.7])
add_heading(doc, "建议携带的兜底材料", 2)
add_bullet(doc, "一张启动成功截图、一张 monitor 截图、一张 GUI/文件管理器截图。")
add_bullet(doc, "make test-release 最终通过的终端截图或日志。")
add_bullet(doc, "本手册的‘一分钟精简路线’和‘故障处理’页面。")

add_heading(doc, "三、推荐演示流程（标准 6–8 分钟）", 1)
add_table(doc, ["时间", "操作", "展示目的", "建议讲解"], [
    ("0:00–0:40", "启动动画 → Shell", "证明自主启动链路", "Loader 播放 VGA 视频帧，恢复文本模式后进入 32 位内核。"),
    ("0:40–1:20", "selftest；status", "快速建立可信度", "先给出自检结果，再用单行状态概览当前 tick、PID 和内存。"),
    ("1:20–2:05", "monitor；Q 返回", "展示抢占式调度", "PIT 驱动任务切换，界面实时统计各任务 ticks 与 CPU 占比。"),
    ("2:05–2:45", "memmap；alloc；memmap；dealloc", "展示页分配与回收", "空闲页下降后恢复，说明位图式物理页管理和可复用性。"),
    ("2:45–3:40", "user；exec demo.oex；ps", "展示 Ring 3 与系统调用", "用户程序不能直接操作 VGA，通过 SYS_WRITE 输出，最后退出码为 42。"),
    ("3:40–4:25", "runfault；ps；selftest", "展示异常隔离", "用户态页故障只终止 PID 4，Shell 和内核继续运行。"),
    ("4:25–5:15", "ls；stat big.txt；disk", "展示 OrangeFS", "可变 extent、跨扇区文件、空闲位图、双目录与数据校验。"),
    ("5:15–6:30", "gui → FILES/APPS；tetris", "展示交互完整度", "桌面、文件 CRUD、文本编辑器和中断驱动小游戏。"),
    ("6:30–7:00", "总结", "回扣主线", "从启动、隔离、存储到应用均由同一内核闭环支撑。"),
], [0.75, 1.45, 1.55, 2.75])

add_heading(doc, "四、逐段讲解话术与观察点", 1)
for title, cmd, say, see in [
    ("1. 启动链路", "make run", "这里不是在宿主系统上运行的普通程序。BIOS 先加载 Boot Sector，Loader 装载内核并切换到保护模式，最终跳转到 0x10000 的 32 位内核入口。", "启动动画结束；出现 32-bit Kernel Started 和 OrangeOS>。"),
    ("2. 自检与状态", "selftest / status", "我先用内核自检建立演示基线。自检覆盖关键描述符、内存与 Ring 3 相关能力；status 用于输出便于回归记录的状态快照。", "SELFTEST PASS；Shell 保持可用。"),
    ("3. 调度与监控", "monitor", "PIT 周期性产生时钟中断，保存当前上下文并轮转任务。Monitor 读取内核维护的累计 ticks，因此看到的是调度结果而不是装饰动画。", "当前 PID 变化；任务状态和 CPU 占比持续刷新。"),
    ("4. 页管理", "memmap / alloc / dealloc", "物理页分配会改变位图统计，释放后页数恢复；用户进程运行时动态占用代码、数据、栈三页，退出后归还。", "used 数值 N→N+1→N；进程路径可展示 N→N+3→N。"),
    ("5. 用户态与 ABI", "user / exec demo.oex / ps", "TSS 提供特权级切换所需的内核栈。用户程序通过 int 0x80 进入内核，SYS_WRITE 会校验特权级、长度与用户地址后先复制到内核缓冲区。", "Ring3 OK；Hello from Ring3 OEX!；code=42。"),
    ("6. 异常隔离", "runfault / ps / selftest", "PID 4 主动访问未映射地址 0x50000000。页故障处理器识别 CPL3，只把进程标记为 FAULTED 并回收资源；若是 Ring 0 页故障才保护性停机。", "PID 4 为 FAULTED；Shell 与 selftest 正常。"),
    ("7. 文件系统", "ls / stat big.txt / disk", "OrangeFS v8 使用目录项、位图和可变 extent。主备目录带版本号与校验和，文件数据也有 16 位校验，读取或执行前都验证完整性。", "big.txt 为 700 字节、占两扇区；disk 显示 free=25/32。"),
    ("8. 界面与应用", "gui / tetris", "GUI 不是独立模拟程序，而是复用内核键盘、定时器和 OrangeFS。文件管理器支持新建、编辑、F2 保存和删除；游戏由 PIT ticks 驱动刷新。", "方向键与 Enter 可操作；Q/Esc 能安全返回 Shell。"),
]:
    add_heading(doc, title, 2)
    add_callout(doc, "命令", cmd)
    p=doc.add_paragraph(); font(p.add_run("讲解："),10.5,True,DARK); font(p.add_run(say),10.5)
    p=doc.add_paragraph(); font(p.add_run("观察："),10.5,True,BLUE); font(p.add_run(see),10.5)

add_heading(doc, "五、时间不足时的一分钟精简路线", 1)
add_number(doc, "selftest：用全项 PASS 证明当前镜像处于可演示基线。")
add_number(doc, "monitor：用 10 秒展示抢占式调度与任务 CPU 占比。")
add_number(doc, "exec demo.oex；ps：展示从磁盘加载 Ring 3 程序、系统调用输出与退出码 42。")
add_number(doc, "gui：进入 FILES 或 APPS，展示系统能力已形成可交互应用。")
add_callout(doc, "收束句", "这四步分别证明了系统的稳定性、调度机制、用户态安全边界和上层交互能力。")

add_heading(doc, "六、可能被问到的问题与参考回答", 1)
questions = [
 ("为什么做这个项目？它的价值是什么？", "这个项目的价值在于把课本中的抽象机制落到一条真实执行链上。每个模块都能通过 QEMU 运行、通过 Shell 观察，并用自动测试验证，因此重点是机制理解与系统整合，而不只是界面展示。"),
 ("哪些部分是你自己实现的？", "从 Boot Sector、Loader、保护模式入口，到 GDT/IDT、PIC/PIT、键盘中断、分页、页分配、调度、Ring 3、系统调用、ATA PIO、OrangeFS、Shell 和 GUI 均在项目汇编源码中实现；QEMU、NASM 和 GNU ld 是构建与运行工具。"),
 ("为什么选择汇编而不是 C？", "项目目标是直观看到 x86 的描述符、寄存器、栈帧和特权级切换。汇编使硬件边界更透明。代价是开发效率和可维护性较低；后续可以保留底层启动与中断入口为汇编，把策略层迁移到 C。"),
 ("从开机到 Shell 的流程是什么？", "BIOS 把启动扇区加载到内存；Boot Sector 加载 Loader；Loader 读取内核，完成显示与运行环境准备，建立保护模式所需结构并跳转；内核初始化 GDT、IDT、PIC/PIT、分页、键盘、内存和文件系统，最后进入 Shell。"),
 ("GDT、IDT、TSS 分别做什么？", "GDT 描述代码段、数据段以及不同特权级；IDT 把中断/异常向量映射到处理入口；TSS 在本项目中主要提供 Ring 3 进入 Ring 0 时的内核栈位置，使系统调用和异常能在可信栈上处理。"),
 ("抢占式调度如何实现？", "PIT 产生周期时钟中断。中断入口保存寄存器上下文，调度器更新任务运行 ticks 和状态，按时间片轮转选择下一任务，再恢复其上下文返回。当前实现规模较小，核心是验证上下文切换闭环。"),
 ("分页和物理页分配有什么区别？", "物理页分配解决‘哪一页可用’，分页解决‘虚拟地址映射到哪一物理页以及是否允许访问’。项目用页位图管理物理页，并为用户代码、数据、栈建立带用户权限的映射。"),
 ("Ring 3 是如何进入和退出的？", "内核准备用户代码、数据和栈映射以及用户态段选择子，通过特权级返回框进入 CPL3；用户程序用 int 0x80 请求服务；系统调用入口验证调用来源和参数，最后通过 SYS_EXIT 回收资源并恢复调度。"),
 ("SYS_WRITE 为什么要复制到内核缓冲区？", "不能直接信任用户指针。内核先验证地址落在当前进程允许的代码、数据或栈映射内，并限制长度，再复制到最长 127 字节的内核缓冲区。这样避免用户越界读取或让内核持续使用可能变化的用户内存。"),
 ("用户态页故障为什么不会拖垮内核？", "页故障处理器结合错误码和当前特权级判断来源。CPL3 故障会记录 CR2 和错误码，把进程标为 FAULTED、回收页并切走；Ring 0 故障说明内核自身失去可信状态，所以选择保护性停机。"),
 ("OrangeFS 的磁盘布局和演进重点是什么？", "当前布局中主目录位于 LBA 82，数据区为 83–114，备份目录位于 115。系统从固定文件扩展到可变 extent、位图分配、目录校验、双目录版本恢复和文件数据校验，重点是逐步增强空间利用率与容错。"),
 ("双目录恢复如何选择有效副本？", "启动时分别检查主目录和备份目录的魔数、校验和与版本号，选择校验有效且版本更新的副本，再同步修复另一份。更新时先写备份、再写主目录，降低中途失败造成两份同时不可用的概率。"),
 ("文件数据损坏时如何处理？", "目录项保存文件内容的 16 位加法校验。cat 和 exec 读取数据后先验证；不一致就返回 File data corrupt，尤其损坏的可执行文件不会进入 Ring 3。"),
 ("为什么自定义 OEX2，而不是 ELF？", "OEX2 是教学型最小格式，只包含入口偏移、代码长度、校验和与机器码，便于在有限内核中完整实现校验、装载和执行。它不具备 ELF 的段、重定位和动态链接能力，后续演进可引入 ELF32。"),
 ("GUI 与游戏能证明什么？", "它们不是系统核心创新，但能证明键盘中断、定时器、显示、文件系统和应用状态机可以协同工作。答辩时应把它们作为内核能力的综合验证，而不是项目主线。"),
 ("如何证明系统不是‘演示脚本’？", "一是现场改变状态：分配页、运行用户进程、触发用户态故障、编辑文件并重启；二是 Makefile 有 16 项 QEMU 自动回归，包含故障注入、资源回收和文件系统恢复，不只检查静态输出。"),
 ("当前项目的局限有哪些？", "主要局限包括单地址空间模型较简化、调度策略固定、文件系统容量和目录槽有限、OEX2 不是通用可执行格式、缺少网络/多核/权限模型，且大部分代码为汇编。回答时主动说明边界，并给出可执行的演进顺序。"),
 ("下一步最值得做什么？", "优先把进程地址空间与内核策略解耦，引入每进程页表和更通用的 ELF32 装载；随后完善系统调用、文件描述符和用户态 Shell；再考虑 C 语言重构、网络和多核。这样每一步都建立在现有 Ring 3 与分页基础上。"),
]
for q,a in questions: add_q(doc,q,a)

add_heading(doc, "七、老师可能继续追问的细节", 1)
add_table(doc, ["追问方向", "回答抓手", "避免的说法"], [
    ("CPU 占比是否精确", "说明按任务累计 ticks 计算，是教学型近似统计", "不要声称等同于现代 OS 性能计数器"),
    ("文件系统是否抗断电", "双目录、写入顺序和校验提升恢复能力", "不要承诺事务级强一致性"),
    ("是否真正多进程", "有独立任务上下文、时钟抢占和 PID 4 生命周期", "不要夸大为完整 POSIX 进程模型"),
    ("安全性如何", "强调 Ring 3、页权限、系统调用参数校验和故障隔离", "不要声称达到生产级安全"),
    ("GUI 是否图形模式", "说明是 VGA 彩色界面与键盘交互，具体模式按源码实现", "不确定时不要虚构窗口系统或鼠标支持"),
    ("测试覆盖如何", "列举 16 项回归与故障注入，说明仍需更细单元测试", "不要用‘完全覆盖’"),
], [1.45, 3.1, 1.95])

add_heading(doc, "八、现场故障处理与切换方案", 1)
add_table(doc, ["现象", "立即处理", "对老师的说明"], [
    ("QEMU 无键盘输入", "点击窗口获取焦点；仍无效则重启已验证镜像", "这是宿主窗口焦点问题，不改变镜像测试结果"),
    ("某命令输错", "Ctrl/退格修正，或重新输入；利用 ↑ 调出历史", "顺带展示命令历史与终端能力"),
    ("界面无法退出", "先按 Esc，再按 Q；必要时重启 QEMU", "说明应用设计了统一返回键"),
    ("持久化结果不保留", "确认不是 -snapshot；切换到预备镜像或截图", "区分快照测试模式和真实写盘模式"),
    ("演示时间被压缩", "立即切换到一分钟精简路线", "保留四个最有证明力的节点"),
    ("系统意外停机", "启动备用镜像，并从 selftest 开始", "先恢复可信基线，再继续关键路径"),
], [1.35, 2.75, 2.4])
add_callout(doc, "红线", "fault 会触发 Ring 0 页故障并停机；shutdown 会直接关闭 QEMU；两者只能在所有主线演示完成后执行。", RED)

add_heading(doc, "九、结束总结话术", 1)
add_callout(doc, "30 秒版本", "OrangeOS 的工作重点是把操作系统的关键机制做成一个可运行、可观察、可回归的闭环：它能从 BIOS 启动进入 32 位内核，通过中断完成抢占式调度和输入，通过分页与 Ring 3 建立安全边界，通过系统调用运行磁盘中的用户程序，并用带校验和冗余目录的文件系统保存数据。GUI、文件管理器和游戏进一步证明这些底层能力可以支撑完整交互。项目目前仍是教学型系统，但核心链路已经打通，也为每进程页表、ELF 装载和用户态生态留下了清晰的演进方向。")

add_heading(doc, "十、答辩前 10 分钟核对单", 1)
for x in [
    "确认当前提交与通过 make test-release 的提交一致。",
    "确认 orange.img 存在且已备份；不要临时 make clean。",
    "QEMU 可启动，键盘焦点正常，OrangeOS> 提示符可见。",
    "依次试运行 selftest、monitor、exec demo.oex、gui，并正常返回。",
    "确认演示计时：标准路线不超过 8 分钟，精简路线不超过 1 分钟。",
    "关闭通知和无关窗口，准备回归截图与关键源码位置。",
    "记住三个数字：16 项发布回归；PID 4 用户进程；OEX2 退出码 42。",
    "记住两个故障地址/容量点：用户故障地址 0x50000000；big.txt 为 700 bytes、2 sectors。",
    "把 fault 和 shutdown 留到最后；非必要不演示内核态停机。",
]: add_bullet(doc,x)

doc.core_properties.title = "OrangeOS v1.0 答辩演示与问答手册"
doc.core_properties.subject = "答辩演示流程、讲解话术、常见问题与故障兜底"
doc.core_properties.author = "OrangeOS 项目组"
doc.save(OUT)
print(OUT)
