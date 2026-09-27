from pathlib import Path
from docx import Document
from build_defense_study_docs import setup, cover, toc, table, chain, bullets, qa, footer

OUT = Path(r"D:\Oranges\docs\OrangeOS答辩学习04-OrangeFS文件系统.docx")


def build():
    d = Document(); setup(d)
    cover(d, "四", "OrangeFS 文件系统", "持久化、完整性校验与目录恢复")
    toc(d, [
        "整体架构与技术路线", "ATA PIO 与 LBA28 扇区读写", "OrangeFS 磁盘布局",
        "目录扇区与目录项", "初始化、挂载与自动恢复", "目录校验与文件校验",
        "文件生命周期：创建、分配、写入、读取与删除", "Shell 命令与答辩演示",
        "与 Ring 3 程序加载的联系", "标准答辩表述", "高频追问与复习清单"
    ])

    d.add_heading("1. 整体架构与技术路线", level=1)
    d.add_paragraph("OrangeFS v8 是一个建立在 ATA 扇区读写之上的教学型文件系统。它不只是把文本放进内存数组，而是把目录元数据和文件内容真正写入 QEMU 使用的磁盘镜像。")
    chain(d, "Shell：ls / cat / touch / write / rm / stat / disk / exec\n→ OrangeFS：查找目录项、空间分配、校验和恢复\n→ ATA PIO：按 LBA 读写 512 B 扇区\n→ orange.img：重启后仍保留数据")
    table(d, ["层次", "职责", "代表代码"], [
        ["Shell", "解析文件命令并显示结果", "kernel/shell32.asm"],
        ["文件系统", "把文件名映射到目录项和数据区", "kernel/filesystem32.asm"],
        ["ATA 驱动", "把 LBA 和缓冲区转换为端口读写", "kernel/ata32.asm"],
        ["磁盘镜像", "保存持久化扇区内容", "orange.img"],
    ])

    d.add_heading("2. ATA PIO 与 LBA28 扇区读写", level=1)
    d.add_paragraph("ATA 驱动使用 Primary Master、PIO 和 LBA28。LBA 是逻辑扇区编号；驱动不理解文件名，只负责把一个 512 B 扇区在磁盘与内存缓冲区之间搬运。")
    table(d, ["端口", "作用"], [["0x1F0", "数据端口"], ["0x1F2", "扇区数量"], ["0x1F3～0x1F5", "LBA 低、中、高字节"], ["0x1F6", "硬盘选择和 LBA 高4位"], ["0x1F7", "命令/状态端口"]])
    d.add_heading("2.1 读取一个扇区", level=2)
    chain(d, "EAX=LBA，EDI=目标缓冲区 → 设置一次传输1扇区 → 命令0x20\n→ 等待 BSY=0、DRQ=1 → rep insw 读取256个word = 512 B")
    d.add_heading("2.2 写入一个扇区", level=2)
    chain(d, "EAX=LBA，ESI=源缓冲区 → 命令0x30 → 等待设备\n→ rep outsw 写256个word → 命令0xE7刷新缓存 → 等待不忙")
    bullets(d, ["驱动使用 pushfd/cli，在 PIO 传输期间暂时关闭中断，结束后恢复原 EFLAGS。", "状态等待有循环上限，但当前接口没有把超时或 ATA ERR 位明确返回上层，这是教学实现的局限。", "代码定位：kernel/ata32.asm。"])

    d.add_heading("3. OrangeFS 的真实磁盘布局", level=1)
    chain(d, "LBA 82：主目录（512 B）\nLBA 83～114：数据池（32扇区，共16 KiB）\nLBA 115：备份目录（512 B）")
    table(d, ["区域", "LBA", "作用"], [["主目录", "82", "主要元数据副本"], ["数据区", "83～114", "32个文件数据扇区"], ["备份目录", "115", "目录损坏时用于恢复"]])
    d.add_paragraph("这是固定布局：没有动态分区分析，容量和目录项数量固定，换来的是结构简单、便于用汇编实现和答辩说明。")

    d.add_heading("4. 目录扇区与目录项", level=1)
    d.add_heading("4.1 一个目录扇区的布局", level=2)
    table(d, ["字节偏移", "内容", "说明"], [["0～3", "魔数 ORF8", "识别 OrangeFS v8"], ["4～7", "32位空间位图", "每一位对应一个数据扇区"], ["8～175", "7个目录项", "每项24 B"], ["176～179", "目录版本号", "判断主备目录的新旧"], ["180～507", "保留区", "后续扩展"], ["508～511", "32位 XOR 校验", "验证整个目录扇区"]])
    d.add_heading("4.2 每个24字节目录项", level=2)
    table(d, ["项内偏移", "字段", "大小"], [["+0～+15", "文件名", "16 B"], ["+16～+17", "文件大小", "2 B"], ["+18～+19", "起始 LBA", "2 B"], ["+20", "占用扇区数", "1 B"], ["+21", "状态标志", "1 B"], ["+22～+23", "文件数据校验", "2 B"]])
    d.add_paragraph("目录项只保存“起始 LBA + 扇区数”，因此一个文件必须占用连续扇区；没有类似 FAT 的数据块链。")

    d.add_heading("5. 空间位图与默认文件", level=1)
    d.add_paragraph("目录偏移4的32位位图对应数据池的32个扇区：0表示空闲，1表示已占用。bit 0 对应 LBA 83，bit 31 对应 LBA 114。默认值 0x7F 表示低7位已占用。")
    table(d, ["默认文件", "起始LBA", "扇区数", "用途"], [["hello.txt", "83", "1", "基础读取演示"], ["about.txt", "84", "1", "版本说明"], ["config.txt", "85", "1", "配置说明"], ["demo.oex", "86", "1", "可执行用户程序"], ["bad.oex", "87", "1", "非法格式测试"], ["big.txt", "88", "2", "700 B跨扇区测试"]])
    d.add_paragraph("默认文件共占7个数据扇区，所以 disk 初始应报告 32−7=25 个空闲扇区。")

    d.add_heading("6. 初始化、挂载与自动恢复", level=1)
    d.add_paragraph("fs_init32 同时读取主目录和备份目录，分别验证魔数与 XOR 校验，然后依据有效性和版本号选择最终目录。")
    chain(d, "读取主目录 LBA82 → 验证 ORF8 + XOR\n读取备份 LBA115 → 验证 ORF8 + XOR\n→ 两份有效：选择版本较新者\n→ 仅一份有效：用有效副本恢复另一份\n→ 两份无效：重新格式化\n→ 把最终目录同步到主、备位置")
    table(d, ["主目录", "备份目录", "处理方式"], [["有效", "有效", "比较版本，采用较新者并同步"], ["有效", "无效", "用主目录修复备份"], ["无效", "有效", "用备份恢复主目录"], ["无效", "无效", "格式化并生成默认文件"]])
    d.add_paragraph("双目录恢复的是目录元数据，不是文件内容。若数据区损坏，系统可以通过文件校验发现，但没有第二份内容用于自动恢复。")

    d.add_heading("7. 两种完整性校验", level=1)
    d.add_heading("7.1 目录：32位 XOR", level=2)
    d.add_paragraph("512 B 目录等于128个 dword。保存时 XOR 前127个 dword，把结果放到最后一个 dword；验证时 XOR 全部128项，结果为0才有效。")
    chain(d, "保存：XOR(dword[0..126]) → 写入 dword[127]\n验证：XOR(dword[0..127]) == 0")
    d.add_heading("7.2 文件：16位累加和", level=2)
    d.add_paragraph("写入时对文件的实际有效字节逐字节累加并截为16位，写入目录项；cat 时重新计算并比较。不一致则返回状态2，Shell 显示 File data corrupt。")
    bullets(d, ["目录校验保护固定长度的元数据结构；文件校验只覆盖 E_SIZE 指定的有效数据。", "简单 XOR 和累加和可能碰撞，不具备密码学安全性；项目目标是展示完整性检测流程。", "“可校验”不等于“可恢复”：恢复仍然需要有效副本或重建策略。"])

    d.add_heading("8. 文件的完整生命周期", level=1)
    d.add_heading("8.1 创建：touch", level=2)
    chain(d, "校验名称非空且少于16 B → 检查是否已存在 → 找空目录项\n→ 写入文件名，size/start/sectors/checksum=0，flags=1 → 保存目录")
    d.add_paragraph("创建空文件时不占用数据扇区，真正写入时再延迟分配。最多只有7个目录项，文件名最多15个可见字符。")
    d.add_heading("8.2 分配空间", level=2)
    d.add_paragraph("长度不超过512 B分配1个空闲位；超过512 B分配两个相邻的空闲位。两扇区掩码从二进制11不断左移，找到位图按位与为0的位置。")
    d.add_paragraph("连续分配实现简单，但可能产生外部碎片：总空闲扇区足够时，也可能找不到两个相邻空位。")
    d.add_heading("8.3 写入：write", level=2)
    chain(d, "查找文件 → 文本复制到1024 B缓冲区（最多1023 B）\n→ 按长度申请1或2个连续扇区 → 先写数据\n→ 更新 size/start/sectors → 计算文件校验\n→ 目录版本+1 → 重算目录校验 → 先写备份，再写主目录")
    d.add_paragraph("当前 resize 实现会在尺寸变化时先释放旧区间再申请新区间；若新区间申请失败，缺少完整回滚。这是答辩中可以主动说明的工程改进点。")
    d.add_heading("8.4 读取：cat", level=2)
    chain(d, "文件名 → fs_find32 → 目录项 start/sectors → 读取数据区\n→ 按 size 重算校验 → 相同则返回内容；不同则报告损坏")
    d.add_heading("8.5 删除：rm", level=2)
    chain(d, "查找目录项 → 清除位图 → 清空目录项关键字段 → 保存双目录")
    d.add_paragraph("删除属于逻辑删除：旧扇区内容不会立即覆盖，只是空间被标记为空闲，之后可被其他文件复用。")

    d.add_heading("9. 跨扇区读写与 big.txt", level=1)
    d.add_paragraph("big.txt 长700 B，占用两个扇区。测试检查偏移0、511、512、699均为 A，偏移700为0。重点检查511/512边界，证明代码确实从第一个扇区连续进入第二个扇区。")
    chain(d, "第1扇区：偏移0～511\n第2扇区：偏移512～699为有效数据，偏移700为字符串结束符")

    d.add_heading("10. Shell 命令与答辩演示", level=1)
    table(d, ["步骤", "命令", "证明内容"], [["查看初始状态", "ls / disk", "目录扫描和位图统计"], ["查看元数据", "stat big.txt", "大小、起始LBA、扇区数"], ["跨扇区读取", "cat big.txt", "多扇区extent读取"], ["创建文件", "touch defense.txt", "新增空目录项"], ["写入并验证", "write defense.txt OrangeOS-defense-demo", "空间分配、持久化和校验"], ["重启验证", "reboot 后 cat defense.txt", "内容写入磁盘镜像"], ["删除回收", "rm defense.txt / disk", "目录清除与位图释放"], ["加载程序", "exec demo.oex", "文件系统连接Ring 3执行链"]])
    d.add_heading("10.1 推荐现场顺序", level=2)
    chain(d, "ls → disk → stat big.txt → cat hello.txt\n→ touch defense.txt → write defense.txt OrangeOS-defense-demo\n→ stat defense.txt → cat defense.txt → reboot → cat defense.txt\n→ disk → rm defense.txt → disk → exec demo.oex")

    d.add_heading("11. 与 Ring 3 程序加载的联系", level=1)
    chain(d, "OrangeFS 持久化 demo.oex → exec 读取并校验 OEX2\n→ 解析头部和代码 → 映射用户代码/数据/栈 → 建立 PID4\n→ iretd 进入 Ring 3 → int 0x80 请求服务 → 退出后回收页面")
    d.add_paragraph("这条链把文件系统、ATA、分页、进程调度、Ring 3 和系统调用连接起来，是项目整体性最强的演示之一。")

    d.add_heading("12. 两分钟标准答辩表述", level=1)
    d.add_paragraph("OrangeFS v8 是建立在 ATA PIO 扇区读写之上的教学型文件系统。底层使用 LBA28 定位磁盘，每次通过 0x1F0～0x1F7 端口传输一个512字节扇区。文件系统采用固定布局：LBA 82保存主目录，LBA 83到114是32扇区数据池，LBA 115保存备份目录。目录扇区包含ORF8魔数、32位空间位图、7个24字节目录项、版本号和目录XOR校验。目录项记录文件名、大小、起始LBA、扇区数、状态以及文件数据校验值。创建空文件时只分配目录项，写入时才从位图中申请一到两个连续扇区；系统先写数据，再计算16位累加校验并更新目录，随后增加版本、重算目录校验，并依次保存备份和主目录。读取时重新计算文件校验，不一致就报告损坏。启动时同时验证主备目录：两份有效则选择版本较新者，一份有效则修复另一份，两份无效则重新格式化。因此系统形成了持久化、校验和目录恢复的基本闭环。它仍然是固定容量、单层目录、无日志的教学设计，但核心文件生命周期已经完整。")

    d.add_heading("13. 高频追问", level=1)
    qa(d, [
        ("ATA 驱动和文件系统有什么区别？", "ATA只认识LBA和512 B扇区；文件系统负责文件名、目录项、空间分配、校验及恢复。"),
        ("为什么需要魔数？", "用于判断磁盘扇区是否符合OrangeFS格式，但仍需配合完整校验防止只看前4字节。"),
        ("为什么需要版本号？", "主备目录都有效但内容不同时，用版本号判断哪一份更新。"),
        ("为什么先写备份再写主目录？", "若中途失败，较新的备份仍可能在重启时通过版本号被选择并恢复主目录。"),
        ("双目录能恢复文件内容吗？", "不能，它只备份目录元数据；文件数据损坏只能检测，当前没有内容副本用于恢复。"),
        ("为什么文件要连续存放？", "目录项只记录起始LBA和扇区数，没有保存块链或索引。"),
        ("为什么最大约1023字节？", "内存文件缓冲区为1024 B，还需要一个字符串结束符；磁盘上最多使用两个扇区。"),
        ("删除后内容真的被擦除了吗？", "没有，只清目录项和位图，属于逻辑删除，原扇区会在后续复用时被覆盖。"),
        ("它是不是日志文件系统？", "不是。双目录、版本和校验提供基础恢复，但没有日志、提交记录和完整事务原子性。"),
        ("如何证明持久化？", "写入文件、重启QEMU、再次cat同一文件；内容仍存在说明写入了orange.img。"),
        ("当前最重要的改进是什么？", "让ATA错误向上传递，并使resize和目录更新具备失败回滚；进一步可加入CRC、日志和非连续块索引。"),
    ])

    d.add_heading("14. 复习清单", level=1)
    bullets(d, [
        "能画出 LBA 82、83～114、115 的磁盘布局。",
        "能说清一个24 B目录项的六个字段。",
        "能解释位图怎样分配一个或两个连续扇区。",
        "能讲清写数据、文件校验、目录版本和双目录保存的先后顺序。",
        "能区分目录可恢复与文件数据仅可检测。",
        "能用 touch→write→cat→reboot→cat→rm 完成持久化演示。",
        "能把 exec demo.oex 与上一部分的 Ring 3 执行链连接起来。",
    ])
    footer(d, "第4部分：OrangeFS文件系统")
    d.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
