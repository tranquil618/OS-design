from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\Oranges")
OUT = ROOT / "docs"
OUT.mkdir(exist_ok=True)

BLUE = "1F4E79"
LIGHT = "DCE6F1"
PALE = "F3F6F9"


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def setup(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.78)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.18
    for name, size in [("Title", 32), ("Subtitle", 15), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)]:
        s = styles[name]
        s.font.name = "Microsoft YaHei"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string(BLUE)
        s.font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].paragraph_format.space_before = Pt(8)
    styles["Heading 2"].paragraph_format.space_after = Pt(3)


def cover(doc, number, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(48)
    r = p.add_run("OrangeOS v1.0 · 答辩学习手册")
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor.from_string("2E75B6")
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(35)
    p.add_run(f"第 {number} 部分\n{title}")
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.add_run(subtitle)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run("代码定位｜执行链路｜答辩表述｜高频追问｜复习清单")
    r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string("5B6573")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(105)
    p.add_run("项目路径：D:\\Oranges\n技术答辩准备资料").font.color.rgb = RGBColor.from_string("6B7280")
    doc.add_page_break()


def toc(doc, items):
    doc.add_heading("阅读导航", level=1)
    doc.add_paragraph("建议先读“总关系”，再顺着执行链理解代码，最后用问答自测。")
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        r = p.add_run(f"{i:02d}  {item}")
        r.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
    doc.add_page_break()


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[j], h, True); shade(t.rows[0].cells[j], BLUE)
        for run in t.rows[0].cells[j].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255,255,255)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            set_cell_text(cells[j], str(val))
            if len(t.rows) % 2 == 1: shade(cells[j], PALE)
    return t


def chain(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Consolas"; r.font.size = Pt(9.5); r.bold = True
    r.font.color.rgb = RGBColor.from_string("17365D")
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), LIGHT); pPr.append(shd)


def bullets(doc, values):
    for value in values:
        doc.add_paragraph(value, style="List Bullet")


def qa(doc, pairs):
    for q, a in pairs:
        p = doc.add_paragraph()
        r = p.add_run("问：" + q); r.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
        p = doc.add_paragraph("答：" + a)
        p.paragraph_format.left_indent = Inches(0.18)


def footer(doc, label):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"OrangeOS 答辩学习手册 · {label}").font.size = Pt(8)


def build_part2(path):
    d = Document(); setup(d)
    cover(d, "二", "x86 核心机制", "GDT、IDT、PIC、PIT、分页与异常处理")
    toc(d, ["一张图理解六个机制", "GDT：保护模式与特权级", "IDT：事件入口表", "PIC 与 PIT：中断路由和时钟", "分页：地址转换与隔离", "异常处理：以页故障为例", "代码中的完整执行链", "标准答辩表述", "高频问题与自测清单"])
    d.add_heading("1. 一张图理解六个机制", level=1)
    d.add_paragraph("这六项不是彼此独立的模块，而是共同回答三个问题：CPU 正在什么权限下运行？事件发生后跳到哪里？哪些地址允许被访问？")
    chain(d, "GDT：身份与权限  →  IDT：事件入口\nPIT 产生时钟 → PIC 编号/转发 → IDT[0x20] → 调度器\n分页检查访问权限 → 非法访问触发 #PF → IDT[14] → 页故障处理")
    table(d, ["机制", "核心职责", "OrangeOS 中的作用"], [
        ["GDT", "定义段和特权级", "Ring 0/Ring 3 代码段、数据段与 TSS"],
        ["IDT", "把向量映射到处理程序", "异常、定时器、键盘和 int 0x80 的入口"],
        ["PIC", "管理外部硬件中断", "重映射 IRQ，并决定开放哪些 IRQ"],
        ["PIT", "周期产生 IRQ0", "约 100 Hz，为抢占调度提供节拍"],
        ["分页", "虚拟地址转换与页级保护", "内核映射及用户页的 U/S、R/W 权限"],
        ["异常处理", "处理 CPU 检测到的错误", "区分 Ring 0 与 Ring 3 页故障并采取不同策略"],
    ])

    d.add_heading("2. GDT：保护模式中的身份与权限", level=1)
    d.add_paragraph("GDT（全局描述符表）保存段描述符。选择子并不是地址本身，而是“表项索引 + TI + RPL”。OrangeOS 通过不同描述符建立内核态与用户态边界。")
    table(d, ["选择子", "含义", "特权级"], [["0x08", "内核代码段", "Ring 0"], ["0x10", "内核数据段", "Ring 0"], ["0x18/0x1B", "用户代码段（后者 RPL=3）", "Ring 3"], ["0x20/0x23", "用户数据段（后者 RPL=3）", "Ring 3"], ["0x28", "TSS", "用于特权级切换"]])
    bullets(d, ["代码定位：kernel/gdt32.asm。", "TSS 中的 SS0=0x10、ESP0 指向内核栈；Ring 3 发生中断时，CPU 据此切换到可信的 Ring 0 栈。", "GDT解决的是段属性和特权级；分页解决的是页级映射与访问权限，两者共同构成保护。"])

    d.add_heading("3. IDT：统一的事件入口表", level=1)
    d.add_paragraph("IDT（中断描述符表）按 0～255 的向量号保存门描述符。异常、硬件中断和软件中断最后都通过它找到入口。")
    table(d, ["向量", "事件", "门权限"], [["0", "除零异常 #DE", "DPL 0"], ["14", "页故障 #PF", "DPL 0"], ["0x20", "PIT/IRQ0", "DPL 0"], ["0x21", "键盘/IRQ1", "DPL 0"], ["0x80", "系统调用", "DPL 3，可由用户态主动调用"]])
    d.add_paragraph("普通门使用属性 0x8E；系统调用门使用 0xEE。关键差别是 DPL=3，使 Ring 3 的 int 0x80 合法，而用户程序不能随意 int 到内核专用入口。代码定位：kernel/idt32.asm。")

    d.add_heading("4. PIC 与 PIT：中断路由和系统时钟", level=1)
    d.add_heading("4.1 PIC", level=2)
    d.add_paragraph("8259A PIC 接收硬件 IRQ、按屏蔽和优先级规则选择请求，再把向量号交给 CPU。OrangeOS 将主片重映射到 0x20～0x27、从片重映射到 0x28～0x2F，避免与 CPU 的 0～31 号异常冲突。")
    d.add_heading("4.2 PIT", level=2)
    d.add_paragraph("PIT 使用约 1.193182 MHz 的输入时钟，装入除数 11931 后约每秒触发 100 次 IRQ0，即约 10 ms 一个 tick。PIC 负责转发，PIT 负责产生节拍。")
    chain(d, "PIT 到点 → IRQ0 → 主 PIC → 向量 0x20 → IDT[0x20]\n→ irq0_timer 保存现场 → ticks++ → EOI → 调度器 → 恢复现场 → iretd")
    d.add_paragraph("EOI 是处理结束通知；若不发送，PIC 可能认为该中断仍在服务中，后续同级中断可能无法继续进入。")

    d.add_heading("5. 分页：地址转换、权限与隔离", level=1)
    d.add_paragraph("32 位两级分页把线性地址拆成页目录索引、页表索引和页内偏移。CR3 指向页目录；设置 CR0.PG 后开启分页。每页 4 KiB。")
    table(d, ["标志", "含义", "项目用法"], [["P", "页面存在", "不存在时访问触发 #PF"], ["R/W", "是否可写", "控制只读/可写"], ["U/S", "用户可访问/仅内核", "内核常用 0x03，用户映射使用 0x07"]])
    bullets(d, ["初始化阶段对受管理内存建立身份映射，使虚拟地址暂时等于物理地址，降低早期内核复杂度。", "用户代码、数据和栈映射必须同时在 PDE 与 PTE 设置 User 位。", "修改映射后用 invlpg 使对应 TLB 项失效，避免 CPU 继续使用旧缓存。", "代码定位：kernel/paging32.asm、kernel/memory32.asm。"])

    d.add_heading("6. 异常处理：以页故障为例", level=1)
    d.add_paragraph("访问不存在或权限不允许的页面时，CPU 触发向量 14。CR2 保存导致故障的线性地址，栈中的错误码说明是否存在、读写方向以及用户/内核来源。")
    chain(d, "非法内存访问 → CPU 产生 #PF(14) → IDT[14] → 读取 CR2/错误码\n→ 检查旧 CS 的低两位：3 表示 Ring 3，0 表示 Ring 0")
    bullets(d, ["Ring 3 故障：记录地址、错误码和 EIP，终止当前用户任务，回收页面并切换到其他任务；隔离故障而不拖垮系统。", "Ring 0 故障：打印诊断信息并停机，因为内核自身已不可信，盲目继续可能扩大破坏。", "异常由 CPU 直接产生，不经过 PIC；PIC 只负责外部硬件 IRQ。"])

    d.add_heading("7. 代码中的三条完整执行链", level=1)
    d.add_heading("7.1 时钟与抢占", level=2); chain(d, "PIT → IRQ0 → PIC(0x20) → IDT → 保存寄存器 → 调度 → 切换 ESP → popad → iretd")
    d.add_heading("7.2 用户非法访问", level=2); chain(d, "Ring 3 访问受保护页 → 分页权限检查失败 → #PF → IDT[14] → 终止/回收当前用户任务")
    d.add_heading("7.3 系统调用入口", level=2); chain(d, "Ring 3 执行 int 0x80 → IDT[0x80](DPL3) → TSS 切到 Ring 0 栈 → 内核检查参数 → 返回用户态")

    d.add_heading("8. 一分钟标准答辩表述", level=1)
    d.add_paragraph("OrangeOS 在进入 32 位保护模式后，首先通过 GDT 建立 Ring 0、Ring 3 段和 TSS，用 IDT 统一登记异常、硬件中断和系统调用入口。PIC 把硬件 IRQ 重映射到 0x20 以后，避免与 CPU 异常冲突；PIT 以约 100 Hz 产生 IRQ0，为系统计时和抢占调度提供节拍。分页由 CR3 指向页目录，利用 PDE/PTE 的 Present、Read/Write 和 User/Supervisor 位实现地址转换与用户/内核隔离。一旦访问越权，CPU 触发页故障，处理程序读取 CR2 和错误码：用户态故障只终止相应任务，内核态故障则打印信息并停机。六个机制合起来形成了权限、事件和内存三方面的基本保护。")

    d.add_heading("9. 高频问题", level=1)
    qa(d, [
        ("GDT 和分页有什么区别？", "GDT提供段属性与特权级，分页提供页级地址映射和权限检查；本项目以平坦分段为主，把细粒度隔离交给分页。"),
        ("IDT 和 PIC 有什么区别？", "PIC处理外部 IRQ 的屏蔽、优先级和向量映射；IDT是 CPU 查找所有中断/异常入口的表。"),
        ("为什么 PIC 要重映射到 0x20？", "0～31 被 CPU 保留给异常，重映射可以避免 IRQ 与异常向量重叠。"),
        ("PIT 为什么设置成约 100 Hz？", "约 10 ms 的粒度足以展示抢占调度，同时不会产生过高的中断开销。"),
        ("CR2、CR3、CR0.PG 分别是什么？", "CR2记录页故障地址，CR3保存页目录物理地址，CR0.PG 控制是否启用分页。"),
        ("为什么页故障不能简单跳过故障指令？", "指令语义可能尚未完成，长度也不一定容易可靠判断；跳过会破坏程序状态。"),
        ("为什么中断入口要保存寄存器？", "被打断任务恢复后必须看到与中断前一致的通用寄存器上下文，调度器也依赖完整现场。"),
        ("为什么用 iretd 而不是 ret？", "iretd 会恢复 EIP、CS、EFLAGS，跨特权级时还恢复 ESP、SS；普通 ret 不具备完整中断返回语义。"),
    ])
    d.add_heading("10. 复习自测", level=1)
    bullets(d, ["能不看资料画出 PIT→PIC→IDT→调度器的链路。", "能解释 0x8E 与 0xEE 的权限差异。", "能说明为什么用户页的 PDE 和 PTE 都要设置 User 位。", "能说出 CR2、CR3 和 CR0.PG 的职责。", "能解释用户页故障与内核页故障为什么采用不同策略。"])
    footer(d, "第2部分：x86核心机制"); d.save(path)


def build_part3(path):
    d = Document(); setup(d)
    cover(d, "三", "抢占调度、Ring 3 与系统调用", "从任务现场切换到受控内核服务")
    toc(d, ["总体技术路线", "任务与 PCB", "PIT 驱动的抢占调度", "Ring 3 用户态如何进入", "int 0x80 系统调用", "用户指针校验", "退出、异常与资源回收", "完整执行链", "标准答辩表述", "高频问题与自测清单"])
    d.add_heading("1. 总体技术路线", level=1)
    d.add_paragraph("这一部分要回答三个连续问题：任务如何被时钟强制切换？用户程序怎样以 Ring 3 权限运行？它又怎样在不获得内核权限的前提下请求服务？")
    chain(d, "PIT 周期中断 → 保存当前任务现场 → 调度器选择下一个 PCB → 恢复现场\n用户任务以 CS=0x1B、SS=0x23 启动 → int 0x80 → TSS 切换内核栈\n→ 内核验证调用号和参数 → 执行有限服务 → iretd 返回 Ring 3")

    d.add_heading("2. 任务模型与 PCB", level=1)
    d.add_paragraph("OrangeOS 使用教学型固定任务表。PCB 记录任务的最小运行状态，真正完整的 CPU 上下文保存在各任务内核栈上的中断帧中。代码定位：kernel/process32.asm。")
    table(d, ["PCB 字段", "大小", "作用"], [["PID", "4 B", "任务编号"], ["STATE", "4 B", "READY/RUNNING/EXITED/FAULTED"], ["ESP", "4 B", "保存该任务内核栈上的现场位置"], ["STACK", "4 B", "任务栈基址或关联信息"]])
    table(d, ["状态", "含义"], [["READY=1", "可被调度"], ["RUNNING=2", "当前正在执行"], ["EXITED=3", "正常结束"], ["FAULTED=4", "因异常终止"]])
    d.add_paragraph("项目中最多维护 4 个任务：PID 1 为 Shell/当前内核任务，PID 2、3 为演示工作任务，PID 4 用于动态用户程序。")

    d.add_heading("3. PIT 驱动的可抢占调度", level=1)
    d.add_paragraph("“可抢占”表示任务不需要主动让出 CPU；PIT 到点后硬件中断会强制进入内核。IRQ0 入口保存通用寄存器，将当前 ESP 交给调度器。")
    chain(d, "IRQ0：pushad → ticks++ → quantum++ → 调度器(current_esp)\n→ 保存 current PCB.ESP → READY → 轮转查找 READY → RUNNING\n→ 返回 next_esp → mov esp,next_esp → popad → iretd")
    bullets(d, ["PIT 约 100 Hz，每 tick 约 10 ms。时间片为 10 tick，因此约 100 ms 切换一次。", "选择策略是循环查找 READY 任务，属于简化的 Round-Robin。", "调度切换的关键不是复制所有寄存器，而是切换 ESP；不同 ESP 指向不同任务已经保存好的现场。", "内核任务初始帧约 44 B：pushad 32 B + EIP/CS/EFLAGS 12 B。用户任务跨级返回还含 ESP/SS，总计约 52 B。"])

    d.add_heading("4. Ring 3 用户态如何进入", level=1)
    d.add_paragraph("用户程序不能仅靠修改变量变成 Ring 3。内核需要准备用户代码页、数据页、用户栈、Ring 3 段选择子以及可被 iretd 消费的完整返回帧。")
    table(d, ["虚拟区域", "地址", "用途"], [["用户代码", "0x40000000", "执行 OEX2 或内置用户代码"], ["用户数据", "0x40001000", "程序数据"], ["用户栈", "0x40002000～0x40002FFF", "栈顶为 0x40003000"]])
    bullets(d, ["usermode_prepare32 分配 3 个物理页，并以用户可访问标志映射到固定虚拟地址。", "process_spawn_user32 构造用户现场：EIP=入口，CS=0x1B，EFLAGS=0x202，ESP=0x40003000，SS=0x23。", "iretd 弹出上述现场后，CPU 的 CPL 变为 3，开始执行用户代码。", "TSS 的 SS0/ESP0 保证用户态发生中断或系统调用时，CPU 先切换到受信任内核栈。"])

    d.add_heading("5. int 0x80：受控系统调用", level=1)
    d.add_paragraph("Ring 3 不能直接访问内核数据或执行特权操作，只能通过 IDT 中 DPL=3 的 0x80 门进入。调用号放在 EAX，其他寄存器传递参数。")
    table(d, ["EAX", "系统调用", "返回/效果"], [["0", "get_ticks", "返回系统 tick"], ["1", "get_free_pages", "返回空闲页数"], ["2", "get_pid", "返回当前 PID"], ["3", "exit", "结束当前用户任务并调度"], ["4", "set_result", "记录用户程序结果"], ["5", "write", "校验并输出用户缓冲区"], ["其他", "unknown", "返回 0xFFFFFFFF"]])
    chain(d, "Ring 3：EAX=调用号，准备参数，int 0x80\n→ IDT[0x80] DPL3 → CPU 切到 TSS.ESP0 → 保存现场\n→ syscall 分派与参数校验 → EAX 放返回值 → iretd")

    d.add_heading("6. 为什么必须校验用户指针", level=1)
    d.add_paragraph("用户程序不可信。SYS_WRITE 不能直接把 ESI 当作内核地址读取，否则用户可让内核访问任意地址，造成信息泄漏、页故障甚至内核崩溃。")
    bullets(d, ["确认调用确实来自 Ring 3。", "长度 ECX 必须非零且不超过 127。", "检查起始地址加长度是否溢出。", "整个缓冲区必须完整落在用户代码页、数据页或栈页的合法范围内。", "先复制到 128 B 内核缓冲区，再由内核打印；这体现 copy-from-user 思想。"])
    d.add_paragraph("答辩重点：系统调用门只解决“允许进入”，参数校验才解决“进入后不能滥用”。")

    d.add_heading("7. 退出、异常与资源回收", level=1)
    d.add_paragraph("用户程序正常 exit 或发生页故障后，都不能继续使用原现场。内核标记任务状态、解除三个用户页映射、释放物理页，再选择一个 READY 任务恢复。")
    chain(d, "正常：SYS_EXIT → EXITED → unmap/free 3 pages → schedule next\n异常：#PF → FAULTED → 记录 CR2/error/EIP → unmap/free 3 pages → schedule next")
    d.add_paragraph("这条路径验证了资源计数应从 N 变为 N+3，再回到 N，说明用户任务生命周期没有持续泄漏。正常退出和故障退出共享回收思路，但状态与诊断信息不同。")

    d.add_heading("8. 一个用户程序的完整执行链", level=1)
    chain(d, "Shell 输入 run/exec → 准备或加载 OEX2 → 分配并映射代码/数据/栈\n→ 构造 PID4 的 52 B 用户现场 → 标记 READY → 调度器选中 PID4\n→ iretd 进入 Ring 3 → 用户代码 int 0x80 请求服务\n→ TSS 切到内核栈 → 校验并执行 syscall → 返回 Ring 3\n→ exit 或 #PF → 回收 3 页 → 切换回 Shell/其他任务")
    d.add_paragraph("其中 user 命令偏向同步演示用户态往返；run 将用户程序作为可调度任务运行；exec 可从文件系统加载 OEX2 可执行内容到 PID4。")

    d.add_heading("9. 一分钟标准答辩表述", level=1)
    d.add_paragraph("OrangeOS 采用 PIT 约 100 Hz 的时钟中断实现抢占。IRQ0 入口先保存寄存器和当前 ESP，调度器每 10 个 tick 按轮转策略从 READY 任务中选择下一个任务，并通过切换 ESP 恢复其现场。用户程序运行前，内核为代码、数据和栈分配三个用户页，并构造包含 EIP、CS、EFLAGS、ESP、SS 的 Ring 3 返回帧，随后通过 iretd 降权执行。用户态不能直接访问内核，只能调用 DPL=3 的 int 0x80 门；CPU 根据 TSS 切换到 Ring 0 内核栈，内核再按 EAX 分派有限系统调用。像 write 这样的调用还会严格校验用户指针和长度。程序正常退出或发生页故障时，内核会标记状态、解除映射、释放页面并调度其他任务，从而实现可抢占、受隔离、可回收的用户程序运行闭环。")

    d.add_heading("10. 高频问题", level=1)
    qa(d, [
        ("协作式和抢占式调度有什么区别？", "协作式依赖任务主动让出 CPU；抢占式由时钟中断强制夺回控制权，单个任务无法长期独占处理器。"),
        ("为什么切换 ESP 就能切换任务？", "每个任务的寄存器现场都压在自己的内核栈上；ESP 决定 popad 和 iretd 从哪一份现场恢复。"),
        ("为什么用户帧比内核帧多 8 B？", "跨特权级 iretd 还要恢复用户 ESP 和 SS，因此多两个 4 B 字段。"),
        ("Ring 3 为什么不能直接调用内核函数？", "直接调用不能建立可信边界，也无法自动切换特权级和内核栈；系统调用门提供受控入口。"),
        ("为什么 int 0x80 门必须 DPL=3？", "软件 int 会检查门的 DPL；DPL 低于当前 CPL 时，Ring 3 主动调用会触发保护异常。"),
        ("TSS 在这里做什么？", "主要提供跨特权级进入内核时使用的 SS0 和 ESP0，不负责本项目的软件调度策略。"),
        ("用户页为什么使用 0x07？", "0x07 表示 Present、Writable、User，使 Ring 3 能访问；同时页目录项也必须允许 User。"),
        ("如何防止 SYS_WRITE 读内核内存？", "检查来源权限、长度、加法溢出和完整地址范围，再复制到内核缓冲区输出。"),
        ("用户程序崩溃为什么不会拖垮系统？", "页保护先阻止越权，#PF 处理器识别 Ring 3 来源，只终止并回收该任务。"),
        ("当前实现有哪些局限？", "固定任务数、固定用户虚拟地址、单地址空间式教学设计、简单轮转，无优先级、阻塞队列和完整进程隔离。"),
    ])
    d.add_heading("11. 复习自测", level=1)
    bullets(d, ["能画出 IRQ0 保存现场、调度、恢复现场的栈变化。", "能解释时间片 10 tick 约等于 100 ms。", "能说清 0x1B、0x23 和 TSS.ESP0 的作用。", "能从 EAX 调用号讲到 syscall 返回值。", "能解释为什么 SYS_WRITE 要先校验再复制。", "能完整讲述 PID4 从创建、运行到退出/故障回收的生命周期。"])
    footer(d, "第3部分：调度、Ring3与系统调用"); d.save(path)


if __name__ == "__main__":
    p2 = OUT / "OrangeOS答辩学习02-x86核心机制.docx"
    p3 = OUT / "OrangeOS答辩学习03-调度Ring3与系统调用.docx"
    build_part2(p2); build_part3(p3)
    print(p2); print(p3)
