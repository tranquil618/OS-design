from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=r"D:\Oranges\docs\OrangeOS完整答辩演示流程.docx"
BLUE="2E74B5"; DARK="1F4D78"; INK="0B2545"; LIGHT="E8EEF5"; PALE="F4F6F9"; GRAY="5B6573"; RED="9B1C1C"

def f(r,size=11,bold=False,color="000000",italic=False,name="Microsoft YaHei"):
    r.font.name=name; rp=r._element.get_or_add_rPr(); rp.rFonts.set(qn("w:eastAsia"),name); rp.rFonts.set(qn("w:ascii"),"Calibri" if name!="Consolas" else name); rp.rFonts.set(qn("w:hAnsi"),"Calibri" if name!="Consolas" else name)
    r.font.size=Pt(size); r.bold=bold; r.italic=italic; r.font.color.rgb=RGBColor.from_string(color); return r
def shade(c,color):
    pr=c._tc.get_or_add_tcPr(); x=OxmlElement("w:shd"); x.set(qn("w:fill"),color); pr.append(x)
def margins(c):
    pr=c._tc.get_or_add_tcPr(); m=OxmlElement("w:tcMar"); pr.append(m)
    for n,v in (("top",90),("bottom",90),("start",120),("end",120)):
        x=OxmlElement("w:"+n); x.set(qn("w:w"),str(v)); x.set(qn("w:type"),"dxa"); m.append(x)
def geom(t,widths):
    t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER; pr=t._tbl.tblPr; tw=pr.first_child_found_in("w:tblW"); tw.set(qn("w:w"),"9360"); tw.set(qn("w:type"),"dxa"); ind=OxmlElement("w:tblInd"); ind.set(qn("w:w"),"120"); ind.set(qn("w:type"),"dxa"); pr.append(ind)
    grid=t._tbl.tblGrid
    for x in list(grid): grid.remove(x)
    for w in widths:
        x=OxmlElement("w:gridCol"); x.set(qn("w:w"),str(int(w*1440))); grid.append(x)
    for row in t.rows:
        for i,c in enumerate(row.cells):
            c.width=Inches(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c); tcw=c._tc.get_or_add_tcPr().first_child_found_in("w:tcW"); tcw.set(qn("w:w"),str(int(widths[i]*1440))); tcw.set(qn("w:type"),"dxa")
def table(doc,heads,rows,widths):
    t=doc.add_table(rows=1,cols=len(heads)); t.style="Table Grid"; rh=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement("w:tblHeader"); rep.set(qn("w:val"),"true"); rh.append(rep)
    for i,x in enumerate(heads): shade(t.rows[0].cells[i],LIGHT); p=t.rows[0].cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0); f(p.add_run(x),10,True,DARK)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0); f(p.add_run(x),9.5)
    geom(t,widths); doc.add_paragraph().paragraph_format.space_after=Pt(0)
def h(doc,text,l=1): doc.add_heading(text,level=l)
def para(doc,text): p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); f(p.add_run(text)); return p
def bullet(doc,text,level=0): p=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2"); p.paragraph_format.space_after=Pt(4); f(p.add_run(text),10.5); return p
def code(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.right_indent=Inches(.25); p.paragraph_format.space_after=Pt(7); pr=p._p.get_or_add_pPr(); s=OxmlElement("w:shd"); s.set(qn("w:fill"),"F2F4F7"); pr.append(s)
    for i,line in enumerate(text.splitlines()):
        if i: p.add_run().add_break()
        f(p.add_run(line),9.2,False,INK,name="Consolas")
def box(doc,label,text,color=DARK):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; shade(t.cell(0,0),PALE); geom(t,[6.5]); p=t.cell(0,0).paragraphs[0]; p.paragraph_format.space_after=Pt(0); f(p.add_run(label+"  "),10.5,True,color); f(p.add_run(text),10.5); doc.add_paragraph().paragraph_format.space_after=Pt(0)
def step(doc,n,title,time,commands,talk,observe=None):
    h(doc,f"{n}. {title}（{time}）",2)
    if commands: code(doc,commands)
    box(doc,"现场话术",talk)
    if observe:
        p=doc.add_paragraph(); f(p.add_run("观察重点："),10.5,True,BLUE); f(p.add_run(observe),10.5)

doc=Document(); s=doc.sections[0]; s.page_width=Inches(8.5); s.page_height=Inches(11); s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1); s.header_distance=s.footer_distance=Inches(.492)
st=doc.styles["Normal"]; st.font.name="Microsoft YaHei"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); st.font.size=Pt(11); st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.25
for n,z,b,a,c in [("Heading 1",16,18,10,BLUE),("Heading 2",13,14,7,BLUE),("Heading 3",12,10,5,DARK)]:
    st=doc.styles[n]; st.font.name="Microsoft YaHei"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); st.font.size=Pt(z); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(c); st.paragraph_format.space_before=Pt(b); st.paragraph_format.space_after=Pt(a); st.paragraph_format.keep_with_next=True
for n in ["List Bullet","List Bullet 2","List Number"]:
    st=doc.styles[n]; st.font.name="Microsoft YaHei"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); st.font.size=Pt(10.5); st.paragraph_format.space_after=Pt(4); st.paragraph_format.line_spacing=1.25
hp=s.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; f(hp.add_run("ORANGEOS v1.0｜答辩演示流程"),8.5,False,GRAY)
fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; f(fp.add_run("演示主线：启动 → 内核 → 用户态 → 文件系统 → GUI"),8.5,False,GRAY)

# workshop-agenda inspired cover
for _ in range(3): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; f(p.add_run("DEFENSE DEMONSTRATION RUNBOOK"),10,True,BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(8); f(p.add_run("OrangeOS 完整答辩演示流程"),26,True,INK)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; f(p.add_run("现场操作顺序、讲解话术、观察点与故障兜底"),14,False,DARK)
table(doc,["建议时长","核心节点","最终目标"],[("约 8 分钟","启动、调度、Ring 3、OrangeFS、GUI","证明系统形成可运行、可观察、可验证的完整闭环")],[1.25,2.65,2.6])
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(70); f(p.add_run("答辩使用版｜2026-08-21"),10,False,GRAY)
doc.add_page_break()

h(doc,"一、演示主线",1)
box(doc,"核心叙事","启动系统 → 内核自检 → 抢占式调度与内存 → Ring 3 用户程序 → 用户态异常隔离 → OrangeFS → GUI 与应用 → 总结。")
para(doc,"整场展示应围绕“完整链路、机制可观察、结果可验证”展开，避免连续输入大量命令却不解释命令背后的操作系统机制。")

h(doc,"二、演示前准备",1)
code(doc,"make test-release")
para(doc,"最后必须出现 OrangeOS release regression PASSED (16/16)。正式演示使用 make run 启动 QEMU。")
for x in ["保留已经通过发布回归的提交与 orange.img 备份。","提前确认 QEMU 键盘焦点、窗口大小和投影清晰度。","不要在答辩前执行 make clean。","持久化演示不能使用 -snapshot。","fault 与 shutdown 只能放在全部主线演示之后。"]: bullet(doc,x)

h(doc,"三、完整演示流程",1)
step(doc,1,"项目介绍","约 40 秒",None,"各位老师好，我的项目是 OrangeOS。它是一个基于 x86 架构、从 Boot Sector 和 Loader 开始自主实现的教学型操作系统。项目实现了 32 位保护模式、中断管理、抢占式任务调度、分页和内存管理、Ring 3 用户态、系统调用，以及支持持久化和故障恢复的 OrangeFS 文件系统。在这些底层机制之上，还实现了 Shell、系统监视器、GUI 文件管理器和小游戏。接下来我会按照“启动、内核机制、用户态、文件系统和交互应用”的顺序进行演示。")
step(doc,2,"系统启动","约 30 秒","make run","系统首先由 BIOS 加载 Boot Sector。Boot Sector 把第二阶段 Loader 加载到内存，Loader 读取内核和启动视频资源，然后开启 A20、建立临时 GDT，进入 32 位保护模式，最后跳转到内核入口。现在看到的 Shell 不是运行在 Windows 或 Linux 上的普通程序，而是直接运行在 OrangeOS 内核环境中。","VGA 启动动画结束后显示 OrangeOS 32-bit Kernel Started!，并出现 OrangeOS> 提示符。")
step(doc,3,"内核自检与状态","约 30 秒","selftest\nstatus","我先运行内核自检，确认描述符、内存管理、分页以及 Ring 3 运行环境处于正常状态。status 会输出当前系统 tick、PID、内存页和调度状态，是适合测试记录的状态快照。","SELFTEST PASS；执行后 Shell 继续正常工作。")
step(doc,4,"抢占式调度","约 40 秒","monitor","OrangeOS 使用 PIT 产生约 100 Hz 的时钟中断。每次中断保存当前任务上下文，由调度器选择下一任务，再恢复对应的寄存器现场。因此这里看到的 PID、运行 tick 和 CPU 占比变化来自真实调度，而不是界面动画。观察约 10 秒后按 Q 返回。","当前 PID、调度切换次数、各任务状态、累计 ticks、CPU 占比和内存页持续变化。")
step(doc,5,"内存分配与回收","约 40 秒","memmap\nalloc\nmemmap\ndealloc\nmemmap","memmap 显示受管物理页的使用数量和可视化占用条。物理页分配器通过位图记录每个页面是否可用。分配后空闲页减少，释放后恢复，说明页面完成了分配、回收和复用闭环。","页数呈 N → N+1 → N。时间充足时可追加 malloc、free、malloc 展示堆块复用。")
step(doc,6,"Ring 3 用户态","约 40 秒","user","这条命令会从 Ring 0 内核态进入 CPL3 用户态。用户程序通过 int 0x80 请求系统服务，内核处理完成后安全返回 Shell。特权级切换时，TSS 提供进入 Ring 0 使用的内核栈。用户代码不能直接访问内核页面和硬件，只能通过系统调用使用内核服务。","输出 Ring3 OK，并返回 OrangeOS>。")
step(doc,7,"从文件系统加载用户程序","约 50 秒","exec demo.oex\nps","demo.oex 不是编译进内核的固定任务，而是存放在 OrangeFS 中的用户程序。内核读取文件，检查 OEX2 格式、长度和校验和，动态分配代码、数据和栈页面，并把它作为 PID 4 在 Ring 3 运行。程序通过 SYS_WRITE 输出文字，最后调用 SYS_EXIT 返回退出码 42。","出现 Hello from Ring3 OEX!；ps 显示 PID 4 state=EXITED、code=42。")
step(doc,8,"用户态异常隔离","约 50 秒","runfault\nps\nselftest","用户进程主动访问未映射地址 0x50000000，触发页故障。处理器根据异常前的特权级判断它来自 Ring 3，只终止 PID 4、记录 CR2 和错误码并回收资源。Shell 和其他任务继续运行。如果故障来自 Ring 0，系统才会保护性停机。","ps 显示 state=FAULTED、fault=0x50000000；随后 selftest 仍然通过。")
step(doc,9,"OrangeFS 文件系统","约 50 秒","ls\nstat big.txt\ndisk","这里显示的是 OrangeFS 的真实目录项，数据通过 ATA PIO 从磁盘镜像读取。big.txt 跨越两个扇区，证明文件系统支持跨扇区读取和可变 extent。OrangeFS 使用 32 位扇区位图管理空间，并通过主备目录、版本号、目录校验和文件数据校验提高恢复能力。","big.txt 显示 size=700 bytes、sectors=2；disk 显示 free=25/32。")
step(doc,10,"持久化文件（可选）","约 50 秒","touch demo\nwrite demo survives\ncat demo\nreboot\ncat demo\nrm demo","重启后仍然可以读取 survives，说明数据已经通过 ATA PIO 写入磁盘镜像，而不只是保存在内存中。本段只能在非快照模式演示；时间紧张时建议跳过，避免重新播放启动动画。","重启前后 cat demo 的内容一致，结束后删除 demo 恢复演示基线。")
step(doc,11,"GUI 与文件管理器","约 1 分钟","gui","GUI 桌面包含 SYSTEM、FILES 和 APPS 三个区域，使用方向键选择、Enter 打开。它复用了内核键盘中断、定时器、VGA 显示和 OrangeFS。FILES 支持选择文件、Enter 打开、输入编辑、F2 保存、N 新建和 Delete 删除。按 Q 或 Esc 返回。","Monitor、文件列表和编辑器均可打开；保存结果可回到 Shell 后通过 cat 验证。")
step(doc,12,"Tetris 应用","约 30 秒","tetris","Tetris 包含七类方块、旋转、碰撞、落地、消行和计分。它不是项目核心机制，但可以综合验证时钟中断、键盘输入、VGA 显示和应用状态管理。操作两三个方块后按 Q 返回即可。","左右键移动、上键旋转、下键软降、空格硬降，Q/Esc 返回 Shell。")

h(doc,"四、结束总结话术",1)
box(doc,"可直接说","以上演示打通了 OrangeOS 的完整运行链路：系统能够从 BIOS 自主启动，进入 32 位保护模式，通过中断完成抢占式调度和键盘输入，通过分页和 Ring 3 建立用户态隔离，通过系统调用运行磁盘中的用户程序，并通过 OrangeFS 实现文件持久化、空间分配、数据校验和目录恢复。GUI、文件管理器和游戏进一步证明这些底层模块能够共同支撑完整的交互应用。当前系统仍是教学型操作系统，下一步可以继续实现每进程独立页表、ELF32 装载、更完整的文件描述符和用户态 Shell。")

h(doc,"五、现场命令清单",1)
h(doc,"标准路线",2)
code(doc,"selftest\nstatus\nmonitor\nmemmap\nalloc\nmemmap\ndealloc\nuser\nexec demo.oex\nps\nrunfault\nps\nselftest\nls\nstat big.txt\ndisk\ngui\ntetris")
h(doc,"一分钟精简路线",2)
code(doc,"selftest\nmonitor\nexec demo.oex\nps\nrunfault\nps\ngui")
para(doc,"精简路线仍能证明系统正常、存在真实调度、支持 Ring 3 用户程序、具有异常隔离，并形成完整交互界面。")

h(doc,"六、现场故障兜底",1)
table(doc,["现象","处理方式","说明"],[("QEMU 不接收输入","点击窗口获取焦点；仍无效则重启备用镜像","通常是宿主窗口焦点问题"),("命令输错","重新输入或按 ↑ 调出历史命令","可顺带展示命令历史"),("全屏应用无法返回","先按 Esc，再按 Q","统一返回键均为 Q/Esc"),("持久化内容消失","确认未使用 -snapshot，切换备用镜像","快照模式不会写回镜像"),("时间不足","立即执行一分钟精简路线","保留最有证明力的四类能力"),("系统意外停机","启动已验证镜像，从 selftest 恢复","先建立可信基线，再继续主线")],[1.45,2.65,2.4])
box(doc,"红线","fault 会触发 Ring 0 页故障并停机，shutdown 会关闭 QEMU；这两个命令只能在全部主线完成后执行。",RED)

doc.core_properties.title="OrangeOS 完整答辩演示流程"; doc.core_properties.subject="现场操作顺序、讲解话术、观察点与故障兜底"; doc.core_properties.author="OrangeOS 项目组"; doc.save(OUT); print(OUT)
