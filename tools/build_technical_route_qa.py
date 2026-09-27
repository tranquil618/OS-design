from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=r"D:\Oranges\docs\OrangeOS技术路线与答辩问题汇总.docx"
BLUE="2E74B5"; DARK="1F4D78"; INK="0B2545"; PALE="F4F6F9"; GRAY="5B6573"

def font(r,size=11,bold=False,color="000000",italic=False,name="Microsoft YaHei"):
    r.font.name=name; rp=r._element.get_or_add_rPr(); rp.rFonts.set(qn("w:eastAsia"),name); rp.rFonts.set(qn("w:ascii"),"Calibri" if name!="Consolas" else name); rp.rFonts.set(qn("w:hAnsi"),"Calibri" if name!="Consolas" else name); r.font.size=Pt(size); r.bold=bold; r.italic=italic; r.font.color.rgb=RGBColor.from_string(color); return r
def h(doc,t,l=1): doc.add_heading(t,level=l)
def p(doc,t): x=doc.add_paragraph(); x.paragraph_format.space_after=Pt(6); font(x.add_run(t)); return x
def bullet(doc,t,level=0): x=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2"); x.paragraph_format.space_after=Pt(4); font(x.add_run(t),10.5); return x
def code(doc,t):
    x=doc.add_paragraph(); x.paragraph_format.left_indent=Inches(.25); x.paragraph_format.right_indent=Inches(.25); x.paragraph_format.space_after=Pt(7); pr=x._p.get_or_add_pPr(); s=OxmlElement("w:shd"); s.set(qn("w:fill"),"F2F4F7"); pr.append(s)
    for i,line in enumerate(t.splitlines()):
        if i: x.add_run().add_break()
        font(x.add_run(line),9.2,False,INK,name="Consolas")
def box(doc,label,text):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; pr=t.cell(0,0)._tc.get_or_add_tcPr(); s=OxmlElement("w:shd"); s.set(qn("w:fill"),PALE); pr.append(s); x=t.cell(0,0).paragraphs[0]; x.paragraph_format.space_after=Pt(0); font(x.add_run(label+"  "),10.5,True,DARK); font(x.add_run(text),10.5); doc.add_paragraph().paragraph_format.space_after=Pt(0)
def qa(doc,n,q,a):
    x=doc.add_paragraph(); x.paragraph_format.space_before=Pt(7); x.paragraph_format.space_after=Pt(3); x.paragraph_format.keep_with_next=True; font(x.add_run(f"{n}. {q}"),11,True,DARK)
    x=doc.add_paragraph(); x.paragraph_format.space_after=Pt(5); font(x.add_run("参考回答："),10.5,True,BLUE); font(x.add_run(a),10.5)

doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1); sec.header_distance=sec.footer_distance=Inches(.492)
st=doc.styles["Normal"]; st.font.name="Microsoft YaHei"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); st.font.size=Pt(11); st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.25
for n,z,b,a,c in [("Heading 1",16,18,10,BLUE),("Heading 2",13,14,7,BLUE),("Heading 3",12,10,5,DARK)]:
    st=doc.styles[n]; st.font.name="Microsoft YaHei"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); st.font.size=Pt(z); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(c); st.paragraph_format.space_before=Pt(b); st.paragraph_format.space_after=Pt(a); st.paragraph_format.keep_with_next=True
for n in ["List Bullet","List Bullet 2"]:
    st=doc.styles[n]; st.font.name="Microsoft YaHei"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); st.font.size=Pt(10.5); st.paragraph_format.space_after=Pt(4); st.paragraph_format.line_spacing=1.25
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(hp.add_run("ORANGEOS v1.0｜技术路线与答辩问答"),8.5,False,GRAY)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(fp.add_run("OrangeOS 答辩准备资料"),8.5,False,GRAY)

for _ in range(4): doc.add_paragraph()
x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(x.add_run("DEFENSE TECHNICAL GUIDE"),10,True,BLUE)
x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.space_before=Pt(12); x.paragraph_format.space_after=Pt(8); font(x.add_run("OrangeOS 技术路线与答辩问题汇总"),25,True,INK)
x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(x.add_run("完整实现路径｜52个高频问题｜回答方法｜关键数字"),14,False,DARK)
x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.space_before=Pt(95); font(x.add_run("答辩准备版｜2026-08-21"),10,False,GRAY)
doc.add_page_break()

h(doc,"一、项目技术路线总览",1)
code(doc,"BIOS\n  ↓\nBoot Sector\n  ↓\nLoader\n  ↓\n32 位保护模式\n  ↓\nGDT / IDT / PIC / PIT\n  ↓\n内存管理与分页\n  ↓\n进程调度\n  ↓\nRing 3 与系统调用\n  ↓\nATA 与 OrangeFS\n  ↓\nShell / Monitor / GUI / Games\n  ↓\n16 项 QEMU 发布回归")
box(doc,"一句话概括","OrangeOS 采用自底向上的技术路线，从启动、保护模式和中断开始，逐步建立内存、调度、Ring 3、系统调用和文件系统，最后通过 Shell、GUI 和自动测试组成一个可运行、可观察、可验证的操作系统闭环。")

routes=[
("1. BIOS 与 Boot Sector","BIOS 将启动扇区加载到物理地址 0x7C00，并验证 0xAA55 签名。Boot Sector 初始化实模式段寄存器，使用 INT 13h 把第二阶段 Loader 读到 0x9000，然后转交控制权。由于启动扇区只有 512 字节，它只承担最小化引导。"),
("2. Loader","Loader 建立实模式栈，通过 BIOS E820 获取物理内存布局，使用 INT 13h Extensions 将 40 KiB 内核加载到 0x10000，播放启动视频，开启 A20，建立临时 GDT，设置 CR0.PE，并通过远跳转进入 32 位保护模式。"),
("3. 内核初始化","内核先关闭中断，建立正式 GDT/TSS 和内核栈，再依次初始化显示、IDT、PIC、PIT、物理页、分页、Ring 3 页面、Heap、OrangeFS 和进程调度，最后执行 sti。初始化顺序按依赖关系确定。"),
("4. GDT、IDT、PIC 与 PIT","GDT定义 Ring 0/Ring 3代码段、数据段和 TSS；IDT安装异常、IRQ和系统调用门；PIC把 IRQ 重映射到 0x20～0x2F；PIT设置为约100 Hz，为系统tick、抢占调度、Monitor和游戏刷新提供时间基准。"),
("5. 物理内存与分页","E820结果用于建立4 KiB物理页位图。分页模块建立页目录和页表、开启CR0.PG，并为用户代码、数据和栈设置用户权限映射。进程退出、kill或页故障后解除映射并归还页面。"),
("6. 内核 Heap","物理页管理解决整页分配，Heap在其上提供kmalloc/kfree，采用基础first-fit策略，满足较小内核对象的动态分配和释放复用。"),
("7. 进程与抢占式调度","PIT中断保存寄存器现场、增加tick、更新任务运行时间并调用调度器。调度器按时间片轮转选择下一任务，通过切换ESP到另一任务的寄存器帧，再由popad和iretd恢复执行。"),
("8. Ring 3 用户态","内核准备用户代码、数据、栈页面及用户段选择子，并由TSS提供特权级切换使用的Ring 0栈。用户程序不能直接操作页表、磁盘和VGA，只能通过受控系统调用请求服务。"),
("9. int 0x80 系统调用","IDT中0x80门的DPL设为3，允许Ring 3调用。当前支持状态查询、SYS_EXIT和SYS_WRITE。SYS_WRITE校验调用来源、用户地址和长度，再复制到最多127字节的内核缓冲区后输出。"),
("10. 用户态异常隔离","runfault让PID 4访问未映射地址0x50000000。页故障处理器读取CR2、错误码和故障前CS；若来自Ring 3，仅标记进程FAULTED并回收资源，若来自Ring 0则保护性停机。"),
("11. ATA PIO","内核通过端口设置LBA和扇区参数、发送读写命令、轮询状态并传输数据，为OrangeFS持久化和用户程序加载提供底层磁盘服务。"),
("12. OrangeFS v8","文件系统使用目录项、32位扇区位图和可变extent。主目录位于LBA 82，数据区为83～114，备份目录位于115。主备目录带版本号和校验，文件内容有16位校验，支持CRUD、跨扇区文件和损坏拒绝。"),
("13. OEX2 用户程序","OEX2包含入口偏移、代码长度、校验和和机器码。exec先验证OrangeFS数据，再验证OEX2格式，动态分配用户页面，创建PID 4，在Ring 3运行，示例通过SYS_WRITE输出并以退出码42结束。"),
("14. Shell、Monitor与GUI","Shell支持命令、历史、Tab补全、长行换行和PageUp/PageDown。Monitor展示PID、内存和任务CPU占比。GUI文件管理器和游戏复用键盘中断、PIT、VGA和OrangeFS，作为底层模块协同工作的综合验证。"),
("15. 测试与发布","make test-release顺序执行16项QEMU回归，覆盖启动、终端、GUI、Ring 3、异常隔离、资源回收、OEX2、跨扇区文件、位图、双目录恢复和数据校验。测试包含动态状态比较与磁盘故障注入。"),
]
for title,text in routes: h(doc,title,2); p(doc,text)

h(doc,"二、答辩高频问题与参考回答",1)
groups={
"A. 项目整体":[
("为什么选择操作系统作为项目？","操作系统能把计算机组成、汇编、数据结构和软件工程联系起来。这个项目把启动、中断、内存、调度、文件系统和用户态等抽象概念落实成一条真实运行链路。"),
("项目的核心成果是什么？","核心成果不是GUI，而是从BIOS启动到Ring 3用户程序执行的完整闭环，并具备分页、系统调用、异常隔离、文件持久化和自动回归验证。"),
("哪些部分是自己实现的？","Boot Sector、Loader、保护模式入口、GDT、IDT、PIC、PIT、键盘中断、分页、内存分配、调度、Ring 3、系统调用、ATA PIO、OrangeFS、Shell、Monitor和GUI均在项目源码中实现；NASM、GNU ld和QEMU是构建与运行工具。"),
("项目最大的难点是什么？","跨特权级切换和异常处理最难，因为同时涉及GDT、TSS、IDT、页权限、CPU压栈格式和调度上下文；任何栈偏移或权限位错误都可能导致三重故障。"),
("项目最大的亮点是什么？","Ring 3程序可从文件系统加载并通过系统调用输出；用户态页故障只终止用户进程；OrangeFS具备双目录恢复和数据校验；关键路径均有QEMU自动测试。")],
"B. 启动与保护模式":[
("BIOS启动后第一步做什么？","BIOS读取启动磁盘第一个扇区到0x7C00，验证末尾0xAA55，然后执行Boot Sector。"),("为什么需要两阶段引导？","Boot Sector只有512字节，只负责加载Loader；Loader再完成内存检测、内核加载、视频播放和保护模式切换。"),("为什么内核放在0x10000？","该位置避开低地址BIOS数据区、引导程序和Loader，并方便用实模式地址1000:0000装载，链接地址也保持一致。"),("为什么开启A20？","不开启时超过1 MiB的地址可能回绕；开启后才能可靠访问更高物理内存。"),("为什么设置CR0.PE后还要远跳转？","远跳转用于重新加载CS并刷新预取队列，使CPU真正使用GDT中的32位代码段。")],
"C. 中断系统":[
("GDT和IDT有什么区别？","GDT描述代码段、数据段、权限级和TSS；IDT描述发生中断或异常时跳转到哪个处理入口。"),("为什么重映射PIC？","BIOS默认IRQ向量与CPU异常向量冲突，映射到0x20之后可将硬件中断与异常分开。"),("为什么发送EOI？","EOI通知PIC当前中断已完成，否则后续同级中断可能无法送达。"),("中断门和陷阱门有什么区别？","中断门进入时自动清IF，陷阱门不清IF；项目使用中断门以简化嵌套处理。"),("为什么int 0x80可以由Ring 3调用？","因为IDT中0x80门的DPL设为3；其他普通门DPL为0，用户程序不能主动调用。")],
"D. 调度与进程":[
("什么是抢占式调度？","任务不需要主动交出CPU，时钟中断会定期打断当前任务，由内核保存上下文并切换。"),("上下文保存什么？","包括通用寄存器、EIP、CS、EFLAGS和栈状态；跨特权级时CPU还保存用户ESP和SS。"),("为什么切换ESP就能切换任务？","每个任务栈里保存完整寄存器现场；切到另一现场的ESP，再恢复寄存器和iretd即可继续该任务。"),("当前调度算法是什么？","时间片轮转，适合小规模教学系统，公平且简单，但没有优先级和动态负载调整。"),("Monitor的CPU占比准确吗？","它按任务累计tick计算，是教学型近似统计，不等同于现代OS的硬件性能计数器。")],
"E. 内存与分页":[
("为什么页面是4 KiB？","4 KiB是32位x86非PAE分页的标准页大小，页表和硬件权限检查均以此为粒度。"),("页面分配和Heap有什么区别？","页面分配管理4 KiB物理页；Heap管理更小的内核对象，减少整页浪费。"),("如何验证用户页面没有泄漏？","进程启动前、运行中、退出后的memmap应为N、N+3、N，自动测试也会比较这三个数。"),("页表如何隔离Ring 3？","用户页面设置User权限，内核页面保持Supervisor权限；CPL3访问内核页会触发页故障。"),("什么是恒等映射？","虚拟地址与物理地址相同，用于简化早期内核启用分页前后的地址连续性。")],
"F. Ring 3与系统调用":[
("Ring 0和Ring 3是什么？","Ring 0是最高权限内核态，Ring 3是受限制用户态；用户程序必须通过系统调用请求内核服务。"),("TSS有什么作用？","主要保存Ring 0的SS0和ESP0，使Ring 3发生中断或系统调用时CPU自动切换到可信内核栈。"),("为什么不能相信用户地址？","用户可能传入未映射、越界或内核区域地址，直接使用会造成崩溃或权限绕过。"),("为什么复制到内核缓冲区？","复制确保内核使用稳定受控的数据，并通过长度上限避免持续访问不可信用户内存。"),("为什么不用sysenter？","int 0x80更直观，便于展示IDT、DPL和特权级切换；sysenter更快但配置与返回路径更复杂。"),("用户态页故障后如何继续？","故障进程不再返回原指令；内核标记FAULTED、回收页面并切换到其他任务的寄存器现场。")],
"G. 文件系统":[
("为什么设计OrangeFS？","为了在可控规模内理解目录、分配、持久化和恢复，先完成完整闭环，而不是直接实现复杂的FAT或ext。"),("什么是extent？","extent是一段连续磁盘扇区，目录项记录起始扇区和长度。"),("为什么使用位图？","位图快速表示每个数据扇区是否已用，创建、扩展和删除时据此分配与回收。"),("为什么先写备份再写主目录？","发生中途失败时，有机会保留旧主目录或新备份；启动时通过校验和版本号选择有效副本。"),("双目录是否绝对抗断电？","不是。它提高恢复能力，但不是事务日志或写时复制，不能保证所有硬件故障下的强一致性。"),("目录校验和文件校验有什么区别？","目录校验保护文件名、位图和extent等元数据；文件校验保护文件内容。"),("为什么单文件最多1023字节？","当前演示布局允许单文件最多两个扇区，并为文本终止符预留空间，这是教学版本限制。")],
"H. OEX2与程序加载":[
("为什么不用ELF？","OEX2结构简单，能在有限代码量中完整实现格式验证、校验、页面分配、装载和执行；ELF的段和重定位更复杂。"),("如何防止损坏程序执行？","先验证OrangeFS文件数据，再检查OEX2魔数、长度和代码校验，任一步失败都拒绝创建进程。"),("退出码42有什么意义？","42只是容易识别的演示值，重点是用户程序可通过SYS_EXIT把结果传回内核并由ps观察。")],
"I. GUI与应用":[
("GUI是核心创新吗？","不是，核心是内核机制；GUI用于综合验证显示、输入、定时器和文件系统协同。"),("GUI有鼠标和窗口系统吗？","当前是VGA彩色桌面和键盘交互，还不是完整窗口系统，也未实现鼠标驱动。"),("游戏能证明什么？","游戏需要定时刷新、实时键盘、显示更新、状态管理和返回Shell，可综合验证多个底层模块。")],
"J. 测试与工程质量":[
("如何证明不是硬编码输出？","现场可改变页面、进程和文件状态；自动测试还会比较动态数值并篡改磁盘进行故障注入。"),("为什么使用QEMU快照？","快照测试不写回基础镜像，保证每次测试从相同状态开始，避免相互污染。"),("16项测试是否完全覆盖？","不是，覆盖主要功能和关键回归，但仍缺少更细单元测试、随机压力和长期运行测试。"),("测试通过是否代表没有Bug？","不能，只能证明已覆盖场景符合预期；未覆盖边界、异常时序和硬件差异仍可能存在问题。")],
"K. 局限与规划":[
("当前最大限制是什么？","地址空间、进程模型、文件系统容量和可执行格式较简化，主要使用汇编也增加维护成本。"),("下一步优先做什么？","优先实现每进程独立页表，其次是ELF32和稳定用户态ABI，再扩展文件描述符与用户态Shell。"),("为什么不先做网络或多核？","网络和多核会引入驱动、同步和并发复杂度，应先完善进程地址空间和系统调用基础。"),("如果重构会怎么做？","保留启动、异常入口和上下文切换等底层汇编，将调度策略、文件系统、内存策略和应用逻辑逐步迁移到C。")]
}
n=1
for group,items in groups.items():
    h(doc,group,2)
    for q,a in items: qa(doc,n,q,a); n+=1

h(doc,"三、回答问题的原则",1)
for title,text in [("不夸大","说明这是教学型操作系统，验证核心机制，但不宣称达到Linux或生产级水平。"),("先结论后机制","先用一句话回答，再解释寄存器、数据结构或执行流程。"),("明确边界","没有实现的功能直接说明，并指出现有架构如何扩展。"),("把界面引回底层","GUI和游戏的价值在于证明键盘、定时器、VGA和文件系统协同。"),("用动态证据","优先引用memmap变化、PID状态、退出码、重启持久化和故障注入结果。")]:
    x=doc.add_paragraph(); x.paragraph_format.space_after=Pt(5); font(x.add_run(title+"："),10.5,True,DARK); font(x.add_run(text),10.5)

h(doc,"四、必须记住的关键数字",1)
for x in ["0x7C00：Boot Sector加载地址","0x9000：Loader加载地址","0x10000：Kernel入口地址","0x20 / 0x21：IRQ0 / IRQ1映射向量","0x80：系统调用向量","约100 Hz：PIT频率","4 KiB：物理页和分页粒度","PID 4：独立用户进程","0x50000000：用户态故障演示地址","127字节：SYS_WRITE内核缓冲区上限","42：demo.oex退出码","LBA 82：OrangeFS主目录","LBA 83～114：OrangeFS数据区","LBA 115：OrangeFS备份目录","16项：发布回归数量"]: bullet(doc,x)

h(doc,"五、最终总结话术",1)
box(doc,"可直接说","OrangeOS采用自底向上的技术路线，从BIOS引导、保护模式和中断系统开始，逐步建立物理内存、分页、调度、Ring 3、系统调用、ATA和OrangeFS，最后用Shell、GUI和自动测试把各模块组成一个完整闭环。它目前仍是教学型系统，但已经能够真实启动、执行用户程序、隔离用户故障、持久化文件并通过16项QEMU回归验证。")

doc.core_properties.title="OrangeOS 技术路线与答辩问题汇总"; doc.core_properties.subject="完整技术路线、52个高频答辩问题与回答方法"; doc.core_properties.author="OrangeOS 项目组"; doc.save(OUT); print(OUT)
